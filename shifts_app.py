import streamlit as st
import calendar
import csv
import io
import base64
import requests

from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="Employee Availability",
    page_icon="📅",
    layout="wide"
)


# ============================================================
# SETTINGS
# ============================================================

DATA_FILE = "availability.csv"


# ============================================================
# SHIFTS
# ============================================================

# Monday-Friday

WEEKDAY_SHIFTS = [
    "03:30-07:30",
    "07:30-15:00",
    "15:00-21:30",
    "21:30-03:30",
]


# Saturday-Sunday

WEEKEND_SHIFTS = [
    "03:30-11:30",
    "11:30-19:30",
    "19:30-03:30",
]


# Public holidays

HOLIDAY_SHIFTS = [
    "03:30-11:30",
    "11:30-19:30",
    "19:30-03:30",
]


# ============================================================
# STREAMLIT SECRETS
# ============================================================

try:

    GITHUB_TOKEN = st.secrets["GITHUB_TOKEN"]
    GITHUB_OWNER = st.secrets["GITHUB_OWNER"]
    GITHUB_REPO = st.secrets["GITHUB_REPO"]

    GITHUB_BRANCH = st.secrets.get(
        "GITHUB_BRANCH",
        "data"
    )

except Exception:

    st.error(
        "GitHub settings are missing from "
        "Streamlit Secrets."
    )

    st.stop()


# ============================================================
# EMPLOYEES
# ============================================================
#
# Example in Streamlit Secrets:
#
# [employees]
# "maria@company.com" = "Maria"
# "john@company.com" = "John"
#
# ============================================================

try:

    EMPLOYEES = dict(
        st.secrets["employees"]
    )

except Exception:

    st.error(
        "Employee list is missing from "
        "Streamlit Secrets."
    )

    st.stop()


# ============================================================
# GITHUB API
# ============================================================

GITHUB_API = (
    "https://api.github.com"
    f"/repos/{GITHUB_OWNER}/{GITHUB_REPO}"
    f"/contents/{DATA_FILE}"
)


def github_headers():

    return {
        "Authorization": (
            f"Bearer {GITHUB_TOKEN}"
        ),
        "Accept": (
            "application/vnd.github+json"
        ),
        "X-GitHub-Api-Version": "2022-11-28"
    }


# ============================================================
# READ CSV FROM GITHUB
# ============================================================

def read_data_from_github():

    response = requests.get(
        GITHUB_API,
        headers=github_headers(),
        params={
            "ref": GITHUB_BRANCH
        },
        timeout=20
    )


    # File does not exist

    if response.status_code == 404:

        return [], None


    if not response.ok:

        raise Exception(
            "GitHub read error: "
            f"{response.status_code} "
            f"{response.text}"
        )


    result = response.json()


    sha = result["sha"]


    content = result["content"]

    content = content.replace(
        "\n",
        ""
    )


    decoded = base64.b64decode(
        content
    ).decode(
        "utf-8"
    )


    reader = csv.DictReader(
        io.StringIO(decoded)
    )


    rows = list(reader)


    return rows, sha


# ============================================================
# SAVE CSV TO GITHUB
# ============================================================

def save_data_to_github(rows, old_sha=None):

    output = io.StringIO()

    fieldnames = [
        "employee",
        "email",
        "work_date",
        "shift"
    ]


    writer = csv.DictWriter(
        output,
        fieldnames=fieldnames
    )


    writer.writeheader()


    for row in rows:

        writer.writerow({
            "employee": row["employee"],
            "email": row["email"],
            "work_date": row["work_date"],
            "shift": row["shift"]
        })


    encoded = base64.b64encode(
        output.getvalue().encode(
            "utf-8"
        )
    ).decode(
        "utf-8"
    )


    payload = {
        "message": (
            "Update employee availability"
        ),
        "content": encoded,
        "branch": GITHUB_BRANCH
    }


    if old_sha:

        payload["sha"] = old_sha


    response = requests.put(
        GITHUB_API,
        headers=github_headers(),
        json=payload,
        timeout=20
    )


    if response.status_code in (
        200,
        201
    ):

        return True


    raise Exception(
        "GitHub save error: "
        f"{response.status_code} "
        f"{response.text}"
    )


# ============================================================
# SAVE EMPLOYEE MONTH
# ============================================================

def save_employee_month(
    employee,
    email,
    year,
    month,
    selected
):

    # Try a few times in case another employee
    # submits at exactly the same time.

    for attempt in range(3):

        rows, sha = read_data_from_github()


        # ----------------------------------------------------
        # Remove existing availability for this employee/month
        # ----------------------------------------------------

        new_rows = []


        for row in rows:

            try:

                row_date = date.fromisoformat(
                    row["work_date"]
                )

            except Exception:

                continue


            if (
                row["email"].lower() == email.lower()
                and row_date.year == year
                and row_date.month == month
            ):

                continue


            new_rows.append(row)


        # ----------------------------------------------------
        # Add new selections
        # ----------------------------------------------------

        for work_date, shift in selected:

            new_rows.append({
                "employee": employee,
                "email": email,
                "work_date": work_date,
                "shift": shift
            })


        try:

            save_data_to_github(
                new_rows,
                sha
            )

            return True

        except Exception as error:

            # GitHub 409 usually means another update
            # happened between our read and write.
            #
            # Try reading the latest version again.

            if "409" in str(error):

                continue

            raise


    return False


# ============================================================
# GET EMPLOYEE AVAILABILITY
# ============================================================

def get_employee_availability(
    email,
    year,
    month
):

    rows, sha = read_data_from_github()


    availability = set()


    for row in rows:

        if row["email"].lower() != email.lower():

            continue


        try:

            row_date = date.fromisoformat(
                row["work_date"]
            )

        except Exception:

            continue


        if (
            row_date.year == year
            and row_date.month == month
        ):

            availability.add(
                (
                    row["work_date"],
                    row["shift"]
                )
            )


    return availability


# ============================================================
# ORTHODOX EASTER
# ============================================================

def calculate_orthodox_easter(year):

    a = year % 4
    b = year % 7
    c = year % 19

    d = (19 * c + 15) % 30
    e = (2 * a + 4 * b - d + 34) % 7

    month = (
        d + e + 114
    ) // 31

    day = (
        (d + e + 114) % 31
    ) + 1


    julian_easter = date(
        year,
        month,
        day
    )


    # Julian -> Gregorian difference

    difference = (
        13 if year < 2100 else 14
    )


    return julian_easter + timedelta(
        days=difference
    )


# ============================================================
# CYPRUS HOLIDAYS
# ============================================================

FIXED_HOLIDAYS = {

    (1, 1):
        "New Year's Day",

    (1, 6):
        "Epiphany",

    (3, 25):
        "Greek National Day",

    (4, 1):
        "Cyprus National Day",

    (5, 1):
        "Labour Day",

    (8, 15):
        "Assumption",

    (10, 1):
        "Cyprus Independence Day",

    (10, 28):
        "Ohi Day",

    (12, 25):
        "Christmas Day",

    (12, 26):
        "Boxing Day",
}


def get_holidays(year):

    holidays = {}


    # Fixed holidays

    for (month, day), name in FIXED_HOLIDAYS.items():

        holidays[
            date(
                year,
                month,
                day
            )
        ] = name


    # Orthodox Easter

    easter = calculate_orthodox_easter(
        year
    )


    holidays[
        easter - timedelta(days=2)
    ] = "Good Friday"


    holidays[
        easter
    ] = "Easter Sunday"


    holidays[
        easter + timedelta(days=1)
    ] = "Easter Monday"


    return holidays


# ============================================================
# SHIFTS FOR DATE
# ============================================================

def get_shifts_for_date(
    work_date,
    holidays
):

    if work_date in holidays:

        return HOLIDAY_SHIFTS


    if work_date.weekday() >= 5:

        return WEEKEND_SHIFTS


    return WEEKDAY_SHIFTS


# ============================================================
# CREATE EXCEL
# ============================================================

def create_excel(rows):

    schedule = {}


    # --------------------------------------------------------
    # Group names
    # --------------------------------------------------------

    for row in rows:

        work_date = row["work_date"]
        shift = row["shift"]
        employee = row["employee"]


        key = (
            f"{work_date} {shift}"
        )


        if key not in schedule:

            schedule[key] = []


        if employee not in schedule[key]:

            schedule[key].append(
                employee
            )


    # --------------------------------------------------------
    # Maximum number of names
    # --------------------------------------------------------

    max_names = 1


    for names in schedule.values():

        max_names = max(
            max_names,
            len(names)
        )


    # --------------------------------------------------------
    # Workbook
    # --------------------------------------------------------

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Schedule"


    # --------------------------------------------------------
    # Header
    # --------------------------------------------------------

    headers = [
        "Date / Shift"
    ]


    for number in range(
        1,
        max_names + 1
    ):

        headers.append(
            f"Name {number}"
        )


    worksheet.append(
        headers
    )


    # --------------------------------------------------------
    # Header styling
    # --------------------------------------------------------

    header_fill = PatternFill(
        start_color="1F4E78",
        end_color="1F4E78",
        fill_type="solid"
    )


    header_font = Font(
        color="FFFFFF",
        bold=True
    )


    for cell in worksheet[1]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    # --------------------------------------------------------
    # Rows
    # --------------------------------------------------------

    for key in sorted(
        schedule.keys()
    ):

        names = schedule[key]

        row = [key]

        row.extend(names)


        while len(row) < (
            max_names + 1
        ):

            row.append("")


        worksheet.append(
            row
        )


    # --------------------------------------------------------
    # Formatting
    # --------------------------------------------------------

    for row in worksheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )


    worksheet.column_dimensions[
        "A"
    ].width = 28


    for column in range(
        2,
        max_names + 2
    ):

        worksheet.column_dimensions[
            get_column_letter(column)
        ].width = 18


    worksheet.freeze_panes = "A2"


    return workbook


# ============================================================
# APP TITLE
# ============================================================

st.title(
    "📅 Employee Availability"
)

st.caption(
    "Select your availability for the month."
)


# ============================================================
# EMPLOYEE
# ============================================================

st.subheader(
    "1. Employee"
)


email = st.text_input(
    "Company email",
    placeholder="name@company.com"
).strip().lower()


employee = None


if email:

    if email in EMPLOYEES:

        employee = EMPLOYEES[email]

        st.success(
            f"Welcome, {employee}!"
        )

    else:

        st.error(
            "This email is not registered."
        )


# ============================================================
# MONTH
# ============================================================

if employee:

    st.subheader(
        "2. Select month"
    )


    today = date.today()


    selected_month = st.date_input(
        "Month",
        value=date(
            today.year,
            today.month,
            1
        ),
        format="DD/MM/YYYY"
    )


    year = selected_month.year

    month = selected_month.month


    month_name = calendar.month_name[
        month
    ]


    holidays = get_holidays(
        year
    )


    # ========================================================
    # LOAD EXISTING
    # ========================================================

    try:

        existing = get_employee_availability(
            email,
            year,
            month
        )

    except Exception as error:

        st.error(
            f"Could not load availability: {error}"
        )

        st.stop()


    # ========================================================
    # CALENDAR
    # ========================================================

    st.subheader(
        f"3. {month_name} {year}"
    )


    st.info(
        "Tick every shift when you are "
        "available to work."
    )


    selected = []


    # --------------------------------------------------------
    # Weekday header
    # --------------------------------------------------------

    weekdays = [
        "MON",
        "TUE",
        "WED",
        "THU",
        "FRI",
        "SAT",
        "SUN"
    ]


    header = st.columns(7)


    for i, weekday in enumerate(
        weekdays
    ):

        with header[i]:

            st.markdown(
                f"**{weekday}**"
            )


    # --------------------------------------------------------
    # Calendar
    # --------------------------------------------------------

    weeks = calendar.monthcalendar(
        year,
        month
    )


    for week_number, week in enumerate(
        weeks
    ):

        columns = st.columns(7)


        for column_number, day_number in enumerate(
            week
        ):

            with columns[column_number]:

                if day_number == 0:

                    st.write("")

                    continue


                work_date = date(
                    year,
                    month,
                    day_number
                )


                # --------------------------------------------
                # Day heading
                # --------------------------------------------

                if work_date in holidays:

                    st.markdown(
                        f"### 🟥 {day_number}"
                    )

                    st.caption(
                        holidays[work_date]
                    )

                elif work_date.weekday() >= 5:

                    st.markdown(
                        f"### 🟦 {day_number}"
                    )

                    st.caption(
                        "Weekend"
                    )

                else:

                    st.markdown(
                        f"### {day_number}"
                    )


                shifts = get_shifts_for_date(
                    work_date,
                    holidays
                )


                # --------------------------------------------
                # Shifts
                # --------------------------------------------

                for shift_number, shift in enumerate(
                    shifts
                ):

                    checkbox_key = (
                        f"{email}_"
                        f"{year}_"
                        f"{month}_"
                        f"{day_number}_"
                        f"{shift_number}"
                    )


                    already_selected = (
                        (
                            work_date.isoformat(),
                            shift
                        )
                        in existing
                    )


                    checked = st.checkbox(
                        shift,
                        value=already_selected,
                        key=checkbox_key
                    )


                    if checked:

                        selected.append(
                            (
                                work_date.isoformat(),
                                shift
                            )
                        )


    # ========================================================
    # SAVE
    # ========================================================

    st.divider()


    if st.button(
        "💾 SAVE AVAILABILITY",
        type="primary",
        use_container_width=True
    ):

        try:

            success = save_employee_month(
                employee,
                email,
                year,
                month,
                selected
            )


            if success:

                st.success(
                    f"Availability saved for "
                    f"{employee} — "
                    f"{month_name} {year}"
                )


                st.rerun()


            else:

                st.error(
                    "The data was changed by "
                    "someone else. Please try again."
                )


        except Exception as error:

            st.error(
                f"Could not save availability: {error}"
            )


# ============================================================
# MANAGER
# ============================================================

st.divider()


st.subheader(
    "🔐 Manager"
)


manager_password = st.text_input(
    "Manager password",
    type="password"
)


try:

    correct_password = st.secrets[
        "MANAGER_PASSWORD"
    ]

except Exception:

    correct_password = ""


if manager_password:

    if manager_password != correct_password:

        st.error(
            "Incorrect manager password."
        )

    else:

        st.success(
            "Manager access enabled."
        )


        # ----------------------------------------------------
        # Manager month
        # ----------------------------------------------------

        today = date.today()


        manager_month = st.date_input(
            "Schedule month",
            value=date(
                today.year,
                today.month,
                1
            ),
            format="DD/MM/YYYY",
            key="manager_month"
        )


        manager_year = (
            manager_month.year
        )

        manager_month_number = (
            manager_month.month
        )


        # ----------------------------------------------------
        # Read all data
        # ----------------------------------------------------

        try:

            all_rows, _ = read_data_from_github()

        except Exception as error:

            st.error(
                f"Could not load data: {error}"
            )

            st.stop()


        # ----------------------------------------------------
        # Filter month
        # ----------------------------------------------------

        rows = []


        for row in all_rows:

            try:

                row_date = date.fromisoformat(
                    row["work_date"]
                )

            except Exception:

                continue


            if (
                row_date.year == manager_year
                and row_date.month == manager_month_number
            ):

                rows.append(
                    row
                )


        # ----------------------------------------------------
        # Schedule
        # ----------------------------------------------------

        if rows:

            schedule = {}


            for row in rows:

                key = (
                    f"{row['work_date']} "
                    f"{row['shift']}"
                )


                if key not in schedule:

                    schedule[key] = []


                if row["employee"] not in schedule[key]:

                    schedule[key].append(
                        row["employee"]
                    )


            max_names = max(
                len(names)
                for names in schedule.values()
            )


            display = []


            for key in sorted(
                schedule.keys()
            ):

                names = schedule[key]


                row = {
                    "Date / Shift": key
                }


                for i in range(
                    max_names
                ):

                    row[
                        f"Name {i + 1}"
                    ] = (
                        names[i]
                        if i < len(names)
                        else ""
                    )


                display.append(
                    row
                )


            st.dataframe(
                display,
                use_container_width=True,
                hide_index=True
            )


            # ------------------------------------------------
            # Excel
            # ------------------------------------------------

            workbook = create_excel(
                rows
            )


            excel_buffer = io.BytesIO()


            workbook.save(
                excel_buffer
            )


            excel_buffer.seek(0)


            st.download_button(
                label="📥 DOWNLOAD EXCEL",
                data=excel_buffer,
                file_name=(
                    f"Employee_Schedule_"
                    f"{calendar.month_name[manager_month_number]}_"
                    f"{manager_year}.xlsx"
                ),
                mime=(
                    "application/vnd.openxmlformats-"
                    "officedocument.spreadsheetml.sheet"
                ),
                use_container_width=True
            )


        else:

            st.info(
                "No availability has been "
                "submitted for this month."
            )
